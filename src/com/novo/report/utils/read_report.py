# -*- coding: UTF-8 -*-
import json
import os
from docx import Document
import sys
import xlwt


def read_doc(file_nm):
    document = Document(file_nm)
    i = 0
    for paragraph in document.paragraphs:
        print("paragraph {}: ".format(i))
        print(paragraph.text)
        i += 1
    i = 0
    for table in document.tables:
        print("table {}:".format(i))
        for row in table.rows:
            print("\t".join([x.text for x in row.cells]))
        i += 1


def read2columns_table(table, result):
    if len(table.rows[0].cells) % 2 != 0:
        return
    for row in table.rows:
        times = int(len(row.cells) / 2)
        for i in range(times):
            result[row.cells[2 * i].text.replace(" ", "").replace("：", "")] = row.cells[2 * i + 1].text


def read_sample_info_table(table, result):
    for row in table.rows:
        for cell in row.cells:
            dd = cell.text.split("：")
            if len(dd) == 2:
                result[dd[0].strip()] = dd[1]


def read_mut_drug_table(table, result):
    drug_info = []
    for row in table.rows[2:]:
        if row.cells[0].text == "/":
            break
        drug_info.append([x.text for x in row.cells])
    result['drug_info'] = drug_info


def read_table(table):
    info = []
    for row in table.rows[1:]:
        if row.cells[0].text == "/":
            break
        info.append([x.text for x in row.cells])
    return info


def read_dmmr_table(table, result):
    for row in table.rows[1:]:
        gene = row.cells[0].text
        result["%s检测结果" % gene] = row.cells[1].text
        result["%s突变丰度" % gene] = row.cells[2].text
        result["%s突变类型" % gene] = row.cells[3].text


def get_var_info(table, result):
    var_type = table.cell(0, 1).text
    var_note = table.cell(3, 1).text
    # print(var_type)
    for c in result['cr_info']:
        # print(var_type1)
        if c[3] in var_type:
            # print(var_note)
            c.append(var_note)
        else:
            c.append('')


def stat_doc(file_nm):
    result = {}
    document = Document(file_nm)
    for table in document.tables:
        if table.cell(0, 0).text.strip() in ["委 托 人：", "检测类型"]:  # 首页表格及检测结果小结
            read2columns_table(table, result)
        elif table.cell(0, 0).text.find("姓名") != -1:  # 样本信息
            read_sample_info_table(table, result)
        elif table.cell(0, 0).text.find("突变基因") != -1:  # 靶向药物总表
            read_mut_drug_table(table, result)
        elif len(table.columns) >= 3 and table.cell(0, 2).text.find("突变类型") != -1:  # 未知临床意义总表
            result['vus_info'] = read_table(table)
        elif len(table.columns) >= 3 and table.cell(0, 2).text.find("外显子") != -1:  # 肿瘤遗传风险
            result['cr_info'] = read_table(table)
        elif table.cell(0, 0).text.find("基因名称") != -1:  # DMMR表格
            read_dmmr_table(table, result)
        elif len(table.rows) == 4 and table.cell(3, 0).text.find("变异解析") != -1:
            # print(table.cell(3, 1).text)
            get_var_info(table, result)
        # elif 有四行，且最后一行第一列是“变异解析”:
        # 	get_var_info(table,result)
    return result


def export_result_to_excel(result=None, title=None, outfile=None):
    wb = xlwt.Workbook()
    ws = wb.add_sheet("stat", cell_overwrite_ok=True)
    # write_merge(行开始, 行结束, 列开始, 列结束, '数据内容')
    ws.write_merge(0, 0, 0, 11, "样本信息")
    ws.write_merge(0, 0, 12, 16, "检测结果小结")
    ws.write_merge(0, 0, 17, 23, "靶向药物用药提示")
    ws.write_merge(0, 0, 24, 35, "错配修复（DMMR）相关基因检测结果")
    ws.write_merge(0, 0, 36, 45, "肿瘤遗传风险检测")
    for index, value in enumerate(title):
        ws.write(1, index, value)
    i = 2
    for rs in result:
        if rs.get('cr_info'):
            col = max(len(rs.get('drug_info')), len(rs.get('cr_info')), 1)
        else:
            col = max(len(rs.get('drug_info')), 1)
        ws.write_merge(i, i + col - 1, 0, 0, rs.get('customer'))
        ws.write_merge(i, i + col - 1, 1, 1, rs.get('product_name'))
        ws.write_merge(i, i + col - 1, 2, 2, rs.get('disease_class_chinese'))
        for index, value in enumerate(title):
            if rs.get(value):
                if col > 1:
                    ws.write_merge(i, i + col - 1, index, index, rs.get(value))
                else:
                    ws.write(i, index, rs.get(value))
        x = title.index("突变基因")
        for index, value in enumerate(rs.get("drug_info", [])):
            for index2, cell in enumerate(value):
                ws.write(i + index, x + index2, cell)
        x = title.index("基因")
        for index, value in enumerate(rs.get("cr_info", [])):
            for index2, cell in enumerate(value):
                ws.write(i + index, x + index2, cell)
        i += col

    wb.save(outfile)


def get_drug_name(did,vd):
    drug_name = did.get('drug_name')
    if did.get('cfda') == 1:
        drug_name += '*'
    if did.get('drug_name') in [d.get('drug_name') for d in vd.get('clinicalList')]:
        drug_name += '#'
    if drug_name:
        drug_name += ', \n'
    return drug_name


def update_json(input_json):
    output_json = {
        '年龄': input_json.get('SampleInfo').get('age') if input_json.get('SampleInfo').get('age') else '未知',
        '姓名': input_json.get('SampleInfo').get('client'),
        '临床诊断': input_json.get('SampleInfo').get('disease_type'),
        '委托人': input_json.get('SampleInfo').get('client'),
        '微卫星不稳定性（MSI）': input_json.get('SummaryOfRresults').get('msi','')+'('+input_json.get('SummaryOfRresults').get('msi_status','')+')',#'0.0（MSS）'
        '报告日期': input_json.get('Analysis').get('created_date').split()[0],#'2020-05-11'
        '样本编号': input_json.get('SampleInfo').get('subbarcode'),
        'disease_class_chinese': input_json.get('Analysis').get('primary_cancer'),
        '标本类型': '组织' if input_json.get('SampleInfo').get('sample_type') == 'tissue' else '血液',
        '送检医院': input_json.get('SampleInfo').get('hospital'),
        '肿瘤遗传风险提示': '共有'+str(input_json.get('SummaryOfRresults').get('crAllListSize','0'))+'个胚系基因突变，其中具有明确或潜在临床意义的基因突变有'+str(input_json.get('SummaryOfRresults').get('hasPathogenicityCount') if input_json.get('SummaryOfRresults').get('hasPathogenicityCount') else input_json.get('SummaryOfRresults').get('crDrugListSize','0'))+'个',
        '肿瘤突变负荷（TMB）': input_json.get('SummaryOfRresults').get('tmb','')+'Muts/Mb('+input_json.get('SummaryOfRresults').get('tmb_status','')+')',
        '样品总体质量评估': '合格',
        '样本条码': input_json.get('SampleInfo').get('subbarcode'),
        'product_name': input_json.get('Analysis').get('product_name'),
        'customer': input_json.get('SampleInfo').get('customer'),
        '委托日期': input_json.get('SampleInfo').get('commission_date'),
        '性别': input_json.get('SampleInfo').get('gender'),
        '靶向用药指导': '共有'+str(input_json.get('SummaryOfRresults').get('thisGeneticmarkerVwListSize','0'))+'个体细胞基因突变，其中具有明确或潜在临床意义的基因突变有'+str(input_json.get('SummaryOfRresults').get('somaticCellMedicationNum','0'))+'个',
    }
    for dmmr in input_json.get('DMMRinfo'):
        output_json[dmmr.get('gene')+'突变类型'] = dmmr.get('mut_type')
        output_json[dmmr.get('gene')+'突变丰度'] = dmmr.get('mutFreq')
        output_json[dmmr.get('gene')+'检测结果'] = dmmr.get('ori_variant')

    cr_info = []
    switch = {'1': '致病性变异',  # 注意此处不要加括号
              '2': '可能致病性变异',  # 注意此处不要加括号
              '3': '不确定性变异',  # 注意此处不要加括号
              '4': '可能良性变异',  # 注意此处不要加括号
              '5': '良性变异',  # 注意此处不要加括号
              '-': '-',
              }
    for cr in input_json.get('CancerRisk',[]):

        one_cr = [cr.get('gene'),cr.get('Chr'),cr.get('Exon'),cr.get('cHGVS'),cr.get('pHGVS'),cr.get('Zygosity'),cr.get('ExonicFunc'),cr.get('c1000g2015aug_all'),switch.get(str(cr.get('rpCr').get('Clinical_significance') or '-')),'']
        if cr.get('rpCr').get('VarClianno') and cr.get('rpCr').get('Clinical_significance') in [1,2]:
            one_cr[-1] = cr.get('rpCr').get('VarClianno')
        cr_info.append(one_cr)
    output_json['cr_info'] = cr_info

    vus_info = []
    drug_info = []
    switch1 = {
        "nonsynonymous SNV": "错义突变",
        "synonymous SNV": "同义突变",
        "nonframeshift insertion": "非移码突变",
        "nonframeshift deletion": "非移码突变",
        "frameshift deletion": "移码突变",
        "frameshift insertion": "移码突变",
        "frameshift indel": "移码突变",
        "nonframeshift indel": "非移码突变",
        "stopgain": "无义突变",
        "stoploss": "stoploss",
        "splicing": "剪切突变",
        "promoter": "启动子区变异",
        "unknown": "未知",
    }
    for vd in input_json.get('VarDrug',[]):
        if vd.get('resultTypeDesc') == '靶向药物':
            di = [
                '共突变' if vd.get('gene')=='Complex' else vd.get('gene'),
                vd.get('ori_variant'),
                '/' if vd.get('mutFreq','.')=='.' else vd.get('mutFreq')+'%',
                '',
                '',
                '',
                '']
            for did in vd.get('drugList',[]):
                drug_name = get_drug_name(did,vd)
                if did.get('approve_range') == '1':
                    di[3] += drug_name
                elif did.get('approve_range') == '2':
                    di[4] += drug_name
                elif did.get('approve_range') == '3':
                    di[5] += drug_name
                elif did.get('approve_range') == '5':
                    di[6] += drug_name
            di[3] = di[3].rstrip(', \n') if di[3].rstrip(', \n') else '无'
            di[4] = di[4].rstrip(', \n') if di[4].rstrip(', \n') else '无'
            di[5] = di[5].rstrip(', \n') if di[5].rstrip(', \n') else '无'
            di[6] = di[6].rstrip(', \n') if di[6].rstrip(', \n') else '无'
            drug_info.append(di)
        elif vd.get('resultTypeDesc') == '未知临床意义' and vd.get('gene') != 'Complex':
            exonic_func = switch1.get(vd.get('ExonicFunc'),vd.get('ExonicFunc'))
            vi = [vd.get('gene'),vd.get('ori_variant'),exonic_func, vd.get('mutFreq')+'%' if vd.get('mutFreq')!='.' else '.']
            vus_info.append(vi)
    output_json['vus_info'] = vus_info
    output_json['drug_info'] = drug_info

    return output_json


def export_result_to_excel_batch():
    title1 = [
        "客户单位",
        "产品名称",
        "匹配癌种",
        "委托人",
        "委托日期",
        "样本编号",
        "报告日期",
        "性别",
        "年龄",
        "临床诊断",
        "送检医院",
        "标本类型",
        "靶向用药指导",
        "肿瘤突变负荷（TMB）",
        "微卫星不稳定性（MSI）",
        "肿瘤遗传风险提示",
        "样品总体质量评估",
        "突变基因",
        "检测结果",
        "突变丰度",
        "可能获益的药物A级",
        "可能获益的药物B级",
        "可能获益的药物C级",
        "可能耐药药物",
        "MLH1检测结果",
        "MLH1突变丰度",
        "MLH1突变类型",
        "MSH2检测结果",
        "MSH2突变丰度",
        "MSH2突变类型",
        "MSH6检测结果",
        "MSH6突变丰度",
        "MSH6突变类型",
        "PMS2检测结果",
        "PMS2突变丰度",
        "PMS2突变类型",
        "基因",
        "染色体",
        "外显子",
        "核苷酸",
        "氨基酸",
        "杂合/纯合",
        "突变类型",
        "千人频率",
        "临床意义",
        "变异解析"
    ]
    results = []
    # f = open('/home/atguigu/read_word/test_report/novopm2_blo_484.json', encoding='utf-8')
    f = open('/home/atguigu/read_word/test0511.json', encoding='utf-8')
    info_json = json.load(f).get('info')
    for i,dd in enumerate(info_json):
        # dd = json.loads(dd)
        result = update_json(dd)
        print(result)
        results.append(result)
    export_result_to_excel(result=results, title=title1, outfile='novopm2_blo_484.xls')
    print('导入完成')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("Usage: python3 {} OutputExcel JsonInfoPath".format(__file__))
    outfile = sys.argv[1]
    results = []
    f = open(sys.argv[2], encoding='utf-8')
    info_json = json.load(f).get('info')
    for dd in info_json:
        dd = json.loads(dd)
        result = update_json(dd)
        results.append(result)
    title = [
        "客户单位",
        "产品名称",
        "匹配癌种",
        "委托人",
        "委托日期",
        "样本编号",
        "报告日期",
        "性别",
        "年龄",
        "临床诊断",
        "送检医院",
        "标本类型",
        "靶向用药指导",
        "肿瘤突变负荷（TMB）",
        "微卫星不稳定性（MSI）",
        "肿瘤遗传风险提示",
        "样品总体质量评估",
        "突变基因",
        "检测结果",
        "突变丰度",
        "可能获益的药物A级",
        "可能获益的药物B级",
        "可能获益的药物C级",
        "可能耐药药物",
        "MLH1检测结果",
        "MLH1突变丰度",
        "MLH1突变类型",
        "MSH2检测结果",
        "MSH2突变丰度",
        "MSH2突变类型",
        "MSH6检测结果",
        "MSH6突变丰度",
        "MSH6突变类型",
        "PMS2检测结果",
        "PMS2突变丰度",
        "PMS2突变类型",
        "基因",
        "染色体",
        "外显子",
        "核苷酸",
        "氨基酸",
        "杂合/纯合",
        "突变类型",
        "千人频率",
        "临床意义",
        "变异解析"
    ]
    export_result_to_excel(result=results, title=title, outfile=outfile)

